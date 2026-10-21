import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_docx_report(doc_a_name, doc_b_name, comparison_text):
    """
    Generates an in-memory Microsoft Word document containing the comparison summary.
    Returns a BytesIO object suitable for downloading via Streamlit.
    """
    doc = Document()

    # 1. Title
    title = doc.add_heading('AI Document Comparison Report', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add metadata (Document Names)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_run = meta.add_run(f"Document A: {doc_a_name}\nDocument B: {doc_b_name}\n")
    p_run.font.size = Pt(10)
    p_run.font.italic = True
    p_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 2. Add Content
    lines = comparison_text.split('\n')
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
            
        # Format headings if the line starts with '#' or matches expected section titles
        if stripped.startswith('#') or any(stripped.startswith(sec) for sec in [
            "Document A Summary", "Document B Summary", "Similarities", 
            "Differences", "Key Details Comparison", "Overall Conclusion", "Limitations"
        ]):
            # Clean heading text
            heading_text = stripped.lstrip('#').strip()
            doc.add_heading(heading_text, level=2)
        else:
            p = doc.add_paragraph(stripped)
            p.paragraph_format.space_after = Pt(4)

    # 3. Save to an in-memory buffer (BytesIO)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer